from pathlib import Path
from zipfile import ZipFile
import re, json
import openpyxl

out=Path(__file__).resolve().parent
source=Path(r'C:\Users\Marc\Documents\MotorHome\Sum of last digits.xlsx')
dest=out/source.name
w=openpyxl.load_workbook(source,data_only=True)
with ZipFile(source) as src, ZipFile(dest,'w') as dst:
    for item in src.infolist():
        content=src.read(item.filename)
        if item.filename=='xl/worksheets/sheet1.xml':
            xml=content.decode('utf-8')
            for row in range(2,26):
                total=sum(w.active.cell(row,col).value for col in range(1,6))
                pattern=rf'(<c\b[^>]*\br="G{row}"[^>]*>)(.*?)(</c>)'
                def change(m):
                    body=re.sub(r'<f\b[^>]*(?:/>|>.*?</f>)',f'<f>SUM(A{row}:E{row})</f>',m[2])
                    body=re.sub(r'<v>.*?</v>',f'<v>{total}</v>',body)
                    return m[1]+body+m[3]
                xml,n=re.subn(pattern,change,xml,flags=re.S)
                assert n==1
            content=xml.encode('utf-8')
        dst.writestr(item,content)
f=openpyxl.load_workbook(dest,data_only=False); v=openpyxl.load_workbook(dest,data_only=True)
for row in range(2,26):
    assert f.active.cell(row,7).value==f'=SUM(A{row}:E{row})'
    assert v.active.cell(row,7).value==sum(v.active.cell(row,c).value for c in range(1,6))
with ZipFile(source) as a, ZipFile(dest) as b:
    changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
    assert changed==['xl/worksheets/sheet1.xml']
original=openpyxl.load_workbook(source,data_only=False)
for row in original.active:
    for c in row:
        new=f.active[c.coordinate]
        if c.column!=7 or c.row==1: assert c.value==new.value
        assert c._style==new._style
print(json.dumps(dict(verified_formulas=24,verified_cached_totals=24,changed_zip_members=changed,output=str(dest))))
