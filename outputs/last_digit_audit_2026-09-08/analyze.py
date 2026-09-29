from pathlib import Path
import json, hashlib, itertools, math
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
SOURCE = Path(r'C:\Users\Marc\Documents\MotorHome\Sum of last digits.xlsx')
ledger = [json.loads(x) for x in (ROOT/'data/draw_history.jsonl').read_text().splitlines() if x.strip()]
w = openpyxl.load_workbook(SOURCE, data_only=True)
wf = openpyxl.load_workbook(SOURCE, data_only=False)
rows = list(w.active.values)[1:]
audit=[]
for rn, row in enumerate(rows,2):
    digits=list(row[:5]); pb=row[5]; g=row[6]
    matches=[r for r in ledger if [x%10 for x in r['main_numbers']]==digits and r['powerball']%10==pb]
    audit.append(dict(excel_row=rn,digits=digits,pb_digit=pb,G=g,main_sum=sum(digits),six_digit_sum=sum(digits)+pb,formula=wf.active.cell(rn,7).value,dates=[r['draw_date'] for r in matches]))
assert all(len(r['dates'])==1 for r in audit), 'Ambiguous workbook chronology'
dates={r['dates'][0] for r in audit}
subset=[r for r in ledger if r['draw_date'] in dates]
assert len(subset)==len(rows)

# Exact legal-line enumeration, preserving without-replacement and slot geometry.
lines=np.array(list(itertools.combinations(range(1,51),5)),dtype=np.uint8)
digits=lines%10
sums=digits.sum(axis=1).astype(int)
counts=np.bincount(sums,minlength=46)
p0=counts/counts.sum()
slot0=np.array([np.bincount(digits[:,j],minlength=10)/len(lines) for j in range(5)])
assert len(lines)==math.comb(50,5) and np.allclose(slot0.sum(axis=0),.5)
delta_null=np.array([np.bincount(np.abs(np.arange(46)-s),weights=p0,minlength=46) for s in range(46)])
def corr(a,lag):
    return float(np.corrcoef(a[:-lag],a[lag:])[0,1])
def analyze(data,seed):
    d=np.array([[n%10 for n in r['main_numbers']] for r in data])
    s=d.sum(axis=1); delta=np.abs(np.diff(s)); n=len(s)
    models={x:[] for x in ['null','sum_frequency','sum_recent8','sum_conditional','adaptive_delta']}
    slot_models={x:[] for x in ['null','frequency','conditional']}
    band=lambda a: np.digitize(a,[18,28])
    for t in range(8,n):
        history=s[:t]
        freq=(np.bincount(history,minlength=46)+20*p0)/(t+20)
        recent=(np.bincount(history[-8:],minlength=46)+20*p0)/28
        prior=s[:t-1]; targets=s[1:t]
        selected=targets[band(prior)==band(s[t-1])]
        conditional=(np.bincount(selected,minlength=46)+20*p0)/(len(selected)+20)
        observed=np.bincount(np.abs(targets-prior),minlength=46)
        expected=delta_null[prior].sum(axis=0)
        avg=expected/len(prior)
        residual=np.divide(observed+20*avg,expected+20*avg,out=np.ones(46),where=expected+20*avg>0)
        adaptive=p0*residual[np.abs(np.arange(46)-s[t-1])]; adaptive/=adaptive.sum()
        for name,p in zip(models,[p0,freq,recent,conditional,adaptive]):
            models[name].append(float(-np.log(p[s[t]])))
        for name in slot_models:
            loss=0
            for j in range(5):
                hist=d[:t,j]
                if name=='null': p=slot0[j]
                else:
                    sample=hist if name=='frequency' else d[1:t,j][d[:t-1,j]==d[t-1,j]]
                    p=(np.bincount(sample,minlength=10)+20*slot0[j])/(len(sample)+20)
                loss-=np.log(p[d[t,j]])
            slot_models[name].append(float(loss/5))
    score={}
    for name,loss in models.items():
        gain=np.array(models['null'])-loss
        score[name]=dict(sum_log_loss=float(np.mean(loss)),line_log_loss=float(np.mean(np.array(loss)+np.log(counts[s[8:]]))),gain_nats=float(gain.mean()),paired_sd=float(gain.std(ddof=1)),positive_targets=int((gain>0).sum()),targets=len(gain))
    # Permutation max tests condition on the observed sum histogram.
    rng=np.random.default_rng(seed)
    perm=np.array([rng.permutation(s) for _ in range(20000)])
    pd=np.abs(np.diff(perm,axis=1))
    obs_count=np.array([np.count_nonzero((delta>=k)&(delta<=k+2)) for k in range(44)])
    perm_count=np.array([((pd>=k)&(pd<=k+2)).sum(axis=1) for k in range(44)]).T
    mean=perm_count.mean(axis=0); sd=perm_count.std(axis=0)
    oz=np.divide(obs_count-mean,sd,out=np.zeros(44),where=sd>0)
    pz=np.divide(perm_count-mean,sd,out=np.zeros_like(perm_count,dtype=float),where=sd>0)
    ac=np.array([corr(s,k) for k in range(1,6)])
    pac=[]
    for k in range(1,6):
        a=perm[:,:-k].astype(float); b=perm[:,k:].astype(float)
        a-=a.mean(axis=1,keepdims=True); b-=b.mean(axis=1,keepdims=True)
        pac.append((a*b).sum(axis=1)/np.sqrt((a*a).sum(axis=1)*(b*b).sum(axis=1)))
    return dict(n=n,start=data[0]['draw_date'],end=data[-1]['draw_date'],digits_counts=np.bincount(d.ravel(),minlength=10).tolist(),sum_sequence=s.tolist(),sum_mean=float(s.mean()),sum_sd=float(s.std(ddof=1)),sum_range=[int(s.min()),int(s.max())],lag_correlations=ac.tolist(),lag_max_permutation_p=float((1+(np.abs(pac).max(axis=0)>=np.abs(ac).max()).sum())/20001),ldsad_11_13=int(((delta>=11)&(delta<=13)).sum()),transitions=n-1,conditional_null_expected_band_hits=float(delta_null[s[:-1],11:14].sum()),max_three_band_permutation_p=float((1+(pz.max(axis=1)>=oz.max()).sum())/20001),best_three_band=[int(oz.argmax()),int(oz.argmax()+2)],score=score,slot_mean_log_loss={k:float(np.mean(v)) for k,v in slot_models.items()})

result=dict(source=str(SOURCE),source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),ledger_sha256=hashlib.sha256((ROOT/'data/draw_history.jsonl').read_bytes()).hexdigest(),audit=audit,missing_ledger_dates=[r['draw_date'] for r in ledger if r['draw_date'] not in dates],null=dict(line_count=len(lines),sum_mean=float(p0@np.arange(46)),sum_sd=float(np.sqrt(p0@((np.arange(46)-22.5)**2))),unconditional_ldsad_11_13=float(p0@delta_null[:,11:14].sum(axis=1)),slot_digit_probabilities=slot0.tolist()),workbook=analyze(subset,20260908),canonical=analyze(ledger,20260909))
(OUT/'results.json').write_text(json.dumps(result,indent=2))
print(json.dumps({k:v for k,v in result.items() if k not in ['audit','null']},indent=2))
print('AUDIT',json.dumps(audit))
print('NULL',json.dumps(result['null']))
